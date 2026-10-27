from __future__ import annotations

import csv
import hashlib
import io
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError

from financeiro.models import (
    Ambito,
    Categoria,
    CompraCartao,
    EventoFinanceiro,
    ItemImportacao,
    NaturezaEvento,
    Pessoa,
    Situacao,
)
from financeiro.services.cartoes_faturas import criar_compra_cartao
from financeiro.services.importacao_ofx import (
    aplicar_sugestao_classificacao,
    preparar_contexto_classificacao,
    sugerir_vinculo_confiavel,
)


CABECALHOS_OBRIGATORIOS = {
    "descricao",
    "data_compra",
    "valor_total",
    "quantidade_parcelas",
}
CABECALHOS_OPCIONAIS = {"categoria", "pessoa", "ambito", "observacao"}


def _normalizar_texto(valor) -> str:
    texto = "" if valor is None else str(valor).strip()
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(caractere for caractere in texto if not unicodedata.combining(caractere)).casefold()


def _valor_json(valor):
    if isinstance(valor, (date, datetime)):
        return valor.isoformat()
    if isinstance(valor, Decimal):
        return str(valor)
    return "" if valor is None else str(valor).strip()


def _cabecalho(valor) -> str:
    return _normalizar_texto(valor).replace(" ", "_").replace("-", "_")


def _ler_csv(conteudo: bytes) -> list[dict]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = conteudo.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValidationError("Não foi possível ler a codificação do CSV.")
    if not texto.strip():
        raise ValidationError("A planilha está vazia.")
    delimitador = ";" if ";" in texto.splitlines()[0] else ","
    leitor = csv.reader(io.StringIO(texto), delimiter=delimitador)
    try:
        cabecalhos = [_cabecalho(valor) for valor in next(leitor)]
    except StopIteration:
        raise ValidationError("A planilha está vazia.")
    return [dict(zip(cabecalhos, linha)) for linha in leitor if any(valor.strip() for valor in linha)]


def _ler_xlsx(conteudo: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError("A leitura XLSX está indisponível no ambiente do aplicativo.") from exc
    try:
        workbook = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("O arquivo XLSX não pôde ser lido.") from exc
    try:
        planilha = workbook.active
        linhas = list(planilha.iter_rows(values_only=True))
        if not linhas:
            return []
        cabecalhos = [_cabecalho(valor) for valor in linhas[0]]
        return [dict(zip(cabecalhos, linha)) for linha in linhas[1:] if any(valor not in (None, "") for valor in linha)]
    finally:
        workbook.close()


def _ler_linhas(arquivo) -> list[dict]:
    arquivo.seek(0)
    conteudo = arquivo.read()
    arquivo.seek(0)
    if arquivo.name.lower().endswith(".xlsx"):
        linhas = _ler_xlsx(conteudo)
    else:
        linhas = _ler_csv(conteudo)
    if not linhas:
        raise ValidationError("A planilha não possui linhas de compras.")
    cabecalhos = set(linhas[0])
    faltantes = sorted(CABECALHOS_OBRIGATORIOS - cabecalhos)
    desconhecidos = sorted(cabecalhos - CABECALHOS_OBRIGATORIOS - CABECALHOS_OPCIONAIS)
    if faltantes:
        raise ValidationError(f"Colunas obrigatórias ausentes: {', '.join(faltantes)}.")
    if desconhecidos:
        raise ValidationError(f"Colunas não reconhecidas: {', '.join(desconhecidos)}.")
    return linhas


def _data(valor, linha: int) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValidationError(f"Linha {linha}: data_compra inválida.")


def _decimal(valor, linha: int) -> Decimal:
    if isinstance(valor, Decimal):
        numero = valor
    else:
        texto = str(valor or "").strip().replace("R$", "").replace(" ", "")
        if "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        try:
            numero = Decimal(texto)
        except (InvalidOperation, ValueError):
            raise ValidationError(f"Linha {linha}: valor_total inválido.")
    numero = numero.quantize(Decimal("0.01"))
    if not numero:
        raise ValidationError(f"Linha {linha}: valor_total deve ser diferente de zero.")
    return numero


def _quantidade(valor, linha: int) -> int:
    try:
        quantidade = int(str(valor).strip())
    except (TypeError, ValueError):
        raise ValidationError(f"Linha {linha}: quantidade_parcelas inválida.")
    if quantidade < 1 or quantidade > 32767:
        raise ValidationError(f"Linha {linha}: quantidade_parcelas deve estar entre 1 e 32767.")
    return quantidade


def _opcoes_por_nome(queryset) -> dict[str, object]:
    opcoes = {}
    for objeto in queryset:
        chave = _normalizar_texto(objeto.nome)
        if chave in opcoes:
            raise ValidationError(f"Há cadastros duplicados para o nome '{objeto.nome}'.")
        opcoes[chave] = objeto
    return opcoes


def ler_compras_planilha(arquivo, usuario) -> list[dict]:
    linhas = _ler_linhas(arquivo)
    categorias = _opcoes_por_nome(Categoria.objects.filter(usuario=usuario, ativa=True))
    pessoas = _opcoes_por_nome(Pessoa.objects.filter(usuario=usuario, ativa=True))
    ambitos = {}
    for valor, rotulo in Ambito.choices:
        ambitos[_normalizar_texto(valor)] = valor
        ambitos[_normalizar_texto(rotulo)] = valor
    resultado = []
    for indice, linha in enumerate(linhas, start=2):
        descricao = str(linha.get("descricao") or "").strip()
        if not descricao:
            raise ValidationError(f"Linha {indice}: descricao é obrigatória.")
        categoria_nome = str(linha.get("categoria") or "").strip()
        pessoa_nome = str(linha.get("pessoa") or "").strip()
        ambito_nome = str(linha.get("ambito") or "").strip()
        categoria = categorias.get(_normalizar_texto(categoria_nome)) if categoria_nome else None
        pessoa = pessoas.get(_normalizar_texto(pessoa_nome)) if pessoa_nome else None
        if categoria_nome and not categoria:
            raise ValidationError(f"Linha {indice}: categoria '{categoria_nome}' não encontrada.")
        if pessoa_nome and not pessoa:
            raise ValidationError(f"Linha {indice}: pessoa '{pessoa_nome}' não encontrada.")
        if ambito_nome and _normalizar_texto(ambito_nome) not in ambitos:
            raise ValidationError(f"Linha {indice}: âmbito '{ambito_nome}' inválido.")
        data_compra = _data(linha.get("data_compra"), indice)
        valor_total = _decimal(linha.get("valor_total"), indice)
        quantidade = _quantidade(linha.get("quantidade_parcelas"), indice)
        resultado.append({
            "linha": indice,
            "descricao": descricao[:220],
            "data_compra": data_compra,
            "valor_total": valor_total,
            "quantidade_parcelas": quantidade,
            "categoria": categoria,
            "pessoa": pessoa,
            "ambito": ambitos.get(_normalizar_texto(ambito_nome), Ambito.NAO_DEFINIDO),
            "ambito_informado": ambitos.get(_normalizar_texto(ambito_nome)) in {
                Ambito.PESSOAL,
                Ambito.PROFISSIONAL,
            },
            "observacao": str(linha.get("observacao") or "").strip(),
            "original": {chave: _valor_json(valor) for chave, valor in linha.items()},
        })
    return resultado


def chave_compra_planejada(usuario_id: int, cartao_id: int, dados: dict) -> str:
    base = "|".join(
        [
            str(usuario_id),
            str(cartao_id),
            _normalizar_texto(dados["descricao"]),
            dados["data_compra"].isoformat(),
            f"{dados['valor_total']:.2f}",
            str(dados["quantidade_parcelas"]),
        ]
    )
    return hashlib.sha256(base.encode("utf-8")).hexdigest()


def _natureza_planilha_cartao(valor: Decimal, descricao: str) -> str:
    descricao_normalizada = _normalizar_texto(descricao).upper()
    if "PAGAMENTO DEBITO EM CONTA" in descricao_normalizada:
        return NaturezaEvento.PAGAMENTO_FATURA
    return NaturezaEvento.ESTORNO if valor < 0 else NaturezaEvento.DESPESA


def criar_itens_planilha(importacao, compras: list[dict]) -> list[ItemImportacao]:
    chaves_existentes = set(
        CompraCartao.objects.filter(
            usuario=importacao.usuario,
            cartao=importacao.cartao,
            chave_compra__gt="",
        ).values_list("chave_compra", flat=True)
    )
    chaves_lote = set()
    itens = []
    for compra in compras:
        chave = chave_compra_planejada(importacao.usuario_id, importacao.cartao_id, compra)
        duplicada = chave in chaves_existentes or chave in chaves_lote
        chaves_lote.add(chave)
        status = ItemImportacao.Status.PENDENTE if duplicada else ItemImportacao.Status.NOVO
        acao = ItemImportacao.Acao.IGNORAR if duplicada else ItemImportacao.Acao.CRIAR
        motivo = (
            "Compra idêntica já planejada; escolha Ignorar ou altere para criar explicitamente."
            if duplicada
            else "Compra futura pronta para revisão."
        )
        item = ItemImportacao(
            importacao=importacao,
            posicao=compra["linha"],
            identificador_externo=f"linha:{compra['linha']}",
            ocorrencia=1,
            chave_conciliacao=chave,
            tipo_normalizado="PLANILHA",
            data=compra["data_compra"],
            valor=compra["valor_total"],
            descricao_original=compra["descricao"],
            natureza=_natureza_planilha_cartao(compra["valor_total"], compra["descricao"]),
            ambito=compra["ambito"],
            numero_parcela=1,
            total_parcelas=compra["quantidade_parcelas"],
            acao=acao,
            observacao=compra["observacao"],
            registro_original={"linha": compra["linha"], **compra["original"]},
            categoria=compra["categoria"],
            pessoa=compra["pessoa"],
            status=status,
            motivo_status=motivo,
        )
        item._categoria_informada = compra["categoria"] is not None
        item._pessoa_informada = compra["pessoa"] is not None
        item._ambito_informado = compra.get("ambito_informado", False)
        itens.append(item)
    for item in itens:
        if item.acao != ItemImportacao.Acao.IGNORAR:
            sugerir_vinculo_confiavel(item)
    contexto = preparar_contexto_classificacao(importacao, itens)
    for item in itens:
        if item.acao != ItemImportacao.Acao.IGNORAR:
            aplicar_sugestao_classificacao(item, contexto)
    ItemImportacao.objects.bulk_create(itens)
    return itens


def criar_compra_planejada(item: ItemImportacao):
    compra = criar_compra_cartao(
        item.importacao.usuario,
        {
            "cartao": item.importacao.cartao,
            "valor": abs(item.valor_final),
            "quantidade_parcelas": item.total_parcelas or 1,
            "ciclo": item.importacao.ciclo,
            "data_original": item.data_final,
            "vencimento": None,
            "natureza": item.natureza or NaturezaEvento.DESPESA,
            "descricao": item.descricao_final,
            "pessoa": item.pessoa,
            "pessoa_nome_livre": "",
            "ambito": item.ambito or Ambito.NAO_DEFINIDO,
            "situacao": Situacao.PREVISTO,
            "categoria": item.categoria,
            "observacao": item.observacao.strip(),
        },
    )
    compra.chave_compra = item.chave_conciliacao
    compra.save(update_fields=["chave_compra", "atualizado_em"])
    parcela = compra.parcelas.get(numero=1)
    return compra, parcela
